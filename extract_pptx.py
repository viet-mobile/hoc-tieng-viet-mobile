import sys
from pptx import Presentation

def extract(path, outpath):
    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n===== SLIDE {i} =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text
                if text.strip():
                    lines.append(text)
            if shape.has_table:
                tbl = shape.table
                for row in tbl.rows:
                    cells = [c.text for c in row.cells]
                    lines.append(" | ".join(cells))
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text
            if note.strip():
                lines.append(f"[NOTE]: {note}")
    with open(outpath, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Wrote {outpath}, {len(prs.slides.__iter__.__self__._sldIdLst)} slides" if False else f"Wrote {outpath}")

if __name__ == "__main__":
    extract(sys.argv[1], sys.argv[2])
