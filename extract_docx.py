import sys
from docx import Document

def extract(path, outpath):
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text for c in row.cells]
            lines.append(" | ".join(cells))
    with open(outpath, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("wrote", outpath, len(lines), "lines")

extract(sys.argv[1], sys.argv[2])
